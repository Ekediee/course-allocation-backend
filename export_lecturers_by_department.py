"""
export_lecturers_by_department.py
─────────────────────────────────
Fetches lecturers and HODs from the /api/v1/users/lecturers endpoint and writes
the results to a formatted Excel (.xlsx) file as a flat table, with a
Department column (departments appear in grouped order).

Usage:
    python export_lecturers_by_department.py
    python export_lecturers_by_department.py --base-url http://localhost:5000 --output my_lecturers.xlsx

Requirements:
    pip install requests openpyxl
"""

import argparse
import sys
from datetime import datetime
from pathlib import Path

try:
    import requests
except ImportError:
    sys.exit("Missing dependency: run  pip install requests")

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Missing dependency: run  pip install openpyxl")


def fetch_lecturers(base_url: str, verify: bool = True, proxies: dict = None) -> list:
    """Call the API and return the list of department groups."""
    url = f"{base_url.rstrip('/')}/api/v1/users/lecturers"
    session = requests.Session()
    if proxies:
        session.proxies.update(proxies)

    try:
        response = session.get(url, timeout=15, verify=verify)
        response.raise_for_status()
    except requests.exceptions.SSLError as e:
        sys.exit(
            f"SSL certificate error: {e}\n"
            "  Tip: if the server uses a self-signed cert, re-run with --no-verify"
        )
    except requests.exceptions.ConnectionError as e:
        sys.exit(
            f"Could not connect to {base_url}\n"
            f"  Detail: {e}\n"
            "  Tips:\n"
            "    • Check that the URL scheme is correct (http vs https)\n"
            "    • If behind a proxy, set the HTTPS_PROXY env var or use --proxy\n"
        )
    except requests.exceptions.Timeout:
        sys.exit(f"Request timed out after 15 seconds connecting to {base_url}.")
    except requests.exceptions.HTTPError:
        sys.exit(f"API error {response.status_code}: {response.text}")
    except requests.exceptions.RequestException as e:
        sys.exit(f"Request failed: {e}")

    return response.json().get("departments", [])


def build_excel(departments: list, output_path: Path) -> None:
    """Write the lecturer data to a formatted, filterable Excel workbook."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lecturers"

    HEADER_FILL  = PatternFill("solid", fgColor="1F3864")   # dark navy
    ALT_ROW_FILL = PatternFill("solid", fgColor="F2F2F2")   # light grey

    HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
    BODY_FONT   = Font(size=10)

    CENTER = Alignment(horizontal="center", vertical="center")
    LEFT   = Alignment(horizontal="left",   vertical="center")

    thin = Side(style="thin", color="BFBFBF")
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

    # (header, API key, width, alignment) - "department" is added from the group name
    COLUMNS = [
        ("Department",             "department",             30, LEFT),
        ("Staff ID",               "staff_id",               14, CENTER),
        ("Name",                   "name",                   30, LEFT),
        ("Role",                   "role",                   10, CENTER),
        ("Gender",                 "gender",                 10, CENTER),
        ("Email",                  "email",                  32, LEFT),
        ("Rank",                   "rank",                   20, LEFT),
        ("Qualification",          "qualification",          18, LEFT),
        ("Specialization",         "specialization",         30, LEFT),
        ("Other Responsibilities", "other_responsibilities", 30, LEFT),
    ]
    last_col = get_column_letter(len(COLUMNS))

    # Title row
    ws.merge_cells(f"A1:{last_col}1")
    title_cell = ws["A1"]
    title_cell.value = f"Lecturers   |   Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    title_cell.font = Font(bold=True, size=12, color="1F3864")
    title_cell.alignment = CENTER
    ws.row_dimensions[1].height = 24

    # Column headers
    for col_idx, (header, _, _, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.alignment = CENTER
        cell.border    = BORDER
    ws.row_dimensions[2].height = 18

    current_row = 3
    total_lecturers = 0
    total_departments = 0

    for dept in departments:
        lecturers = dept.get("lecturers", [])
        if not lecturers:
            continue
        total_departments += 1
        dept_name = dept.get("name", "Unknown Department")

        for lecturer in lecturers:
            row_fill = ALT_ROW_FILL if total_lecturers % 2 == 1 else None
            record = {**lecturer, "department": dept_name}

            for col_idx, (_, key, _, align) in enumerate(COLUMNS, start=1):
                cell = ws.cell(row=current_row, column=col_idx, value=record.get(key) or "")
                cell.font      = BODY_FONT
                cell.alignment = align
                cell.border    = BORDER
                if row_fill:
                    cell.fill = row_fill

            ws.row_dimensions[current_row].height = 15
            current_row += 1
            total_lecturers += 1

    for col_idx, (_, _, width, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{last_col}{current_row - 1}"

    # Summary footer
    current_row += 1
    ws.merge_cells(f"A{current_row}:{last_col}{current_row}")
    summary_cell = ws.cell(
        row=current_row, column=1,
        value=f"Total: {total_departments} department(s),  {total_lecturers} lecturer(s)"
    )
    summary_cell.font      = Font(italic=True, size=9, color="666666")
    summary_cell.alignment = LEFT

    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(
        description="Export lecturers grouped by department from the API to an Excel file."
    )
    parser.add_argument(
        "--base-url",
        default="http://localhost:5000",
        help="Base URL of the running Flask API (default: http://localhost:5000)"
    )
    parser.add_argument(
        "--output",
        default=None,
        help="Output Excel filename (default: lecturers_<timestamp>.xlsx)"
    )
    parser.add_argument(
        "--no-verify",
        action="store_true",
        default=False,
        help="Disable SSL certificate verification (use if the server has a self-signed cert)"
    )
    parser.add_argument(
        "--proxy",
        default=None,
        help="Explicit proxy URL, e.g. http://proxy.university.edu:8080"
    )

    args = parser.parse_args()

    proxies = {"http": args.proxy, "https": args.proxy} if args.proxy else None
    verify  = not args.no_verify

    output_filename = args.output or f"lecturers_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    output_path = Path(output_filename)

    print(f"Fetching lecturers from {args.base_url} ...")
    if not verify:
        print("  ⚠  SSL verification disabled.")
    if proxies:
        print(f"  Using proxy: {args.proxy}")

    departments = fetch_lecturers(args.base_url, verify=verify, proxies=proxies)

    if not departments:
        print("No lecturers found. Excel file will not be created.")
        sys.exit(0)

    total = sum(len(d.get("lecturers", [])) for d in departments)
    print(f"  -> {len(departments)} department(s), {total} lecturer(s) found.")

    print(f"Writing Excel file: {output_path} ...")
    build_excel(departments, output_path)

    print(f"Done! File saved: {output_path.resolve()}")


if __name__ == "__main__":
    main()
